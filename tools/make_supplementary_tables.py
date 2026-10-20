# -*- coding: utf-8 -*-
"""
make_supplementary_tables.py
生成投稿用补充表 S1-S3：新兴 PFAS 替代物候选清单（含身份核验与证据等级列）

输入 : results_alternatives/emerging_alternatives.csv   (2,390 条 read-across 候选)
输出 : supplementary_tables/
         Table_S1_Tier_I_priority_candidates.csv    (100 条)
         Table_S2_Regrettable_substitutions.csv     (239 条)
         Table_S3_All_candidates.csv                (2,390 条)
         Supplementary_Tables_S1-S3.xlsx            (README + S1 + S2 + S3)

行数、工作表名与 README 中的计数一律由数据现算，不写死 —— 历史上这些常量
先后停留在 202 / 2,048，与表体实际的 239 / 2,390 不符。

命名注意：四个结构标记列为 Marker_*，命中数为 MarkerCount_0_4。**不要写回
PMT_** 字样，也不得引用旧的 PMT_* 列名。

证据等级仅基于清单元数据（DSSTox 收录 / CAS 登记 / ToxCast-Tox21 体外记录），
不构成"该物质已在环境中检出"的证明。文献、专利、环境检出三个维度未自动核验，
核验路径见 README 工作表。
"""
import os
import numpy as np
import pandas as pd

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))   # code/
ROOT = os.path.dirname(HERE)                                          # 项目根
SRC = os.path.join(HERE, "results_alternatives", "emerging_alternatives.csv")
OUT = os.path.join(ROOT, "supplementary_tables")
os.makedirs(OUT, exist_ok=True)

# ---------------------------------------------------------------- 中英映射
FAMILY_MAP = {
    "醚型含氟前体/中间体": "Ether-type fluorinated precursor/intermediate",
    "醚型 PFAA（羧酸）": "Ether-type PFAA (carboxylic acid)",
    "磺酰胺/磺酰乙酸类": "Sulfonamide / sulfonyl acetic acid",
    "短链 PFCA": "Short-chain PFCA",
    "醚型 PFAA（磺酸）": "Ether-type PFAA (sulfonic acid)",
    "其他可电离含氟物": "Other ionisable fluorinated substance",
    "短链 PFSA": "Short-chain PFSA",
    "氟磺酰/氟硫酸酯（高反应性中间体）": "Fluorosulfonyl / fluorosulfate (highly reactive intermediate)",
    # screen_alternatives.py 对端基为 CH2CH2SO3H 的调聚物单独给族名；此前本表缺这一条，
    # 该值经下面的 .fillna(...) 原样透传，导致 S3 第 547 行出现中文族名。
    "氟调聚物磺酸（FTS）": "Fluorotelomer sulfonate",
}
FG_MAP = {
    "酯": "Ester",
    "羧酸": "Carboxylic acid",
    "磺酰胺": "Sulfonamide",
    "磺酸": "Sulfonic acid",
    "磺酰/砜": "Sulfonyl / sulfone",
    "其他": "Other",
    "氟磺酰基中间体": "Fluorosulfonyl intermediate",
    "醇": "Alcohol",
    "酰胺": "Amide",
}
CLASS_MAP = {
    "PFCAs": "PFCAs",
    "PFSAs": "PFSAs",
    "磺酰胺/磺酰类": "Sulfonamides / sulfonyls",
}
SPACE_MAP = {
    "B: PFAS 替代物化学空间": "B: PFAS replacement chemical space",
}
TIER_MAP = {
    "I: 优先关注": "I (priority)",
    "II: 值得跟踪": "II (watch)",
    "III: 储备观察": "III (reserve)",
}


def yn(series):
    return np.where(series.astype(str).str.strip().str.lower().isin(["true", "1", "yes"]), "Yes", "No")


def translate(series, mapping, column):
    """严格映射：未登记的取值直接中止，不做静默透传。

    此前五个注释列都用 ``.map(MAP).fillna(raw)``，任何未登记的取值会原样进入交付表。
    中文族名「氟调聚物磺酸（FTS）」就是这样穿过整条导出链、进入英文补充表的
    （S3 第 547 行）。交付件是英文的，透传即缺陷，所以这里把缺失翻译当作错误。
    """
    got = series.map(mapping)
    missing = sorted(set(series[got.isna()].astype(str)))
    if missing:
        raise SystemExit(
            f"ABORT [{column}]: {len(missing)} unmapped value(s) would leak verbatim "
            f"into the English supplementary table: {missing}\n"
            f"        add them to the corresponding *_MAP dict"
        )
    out = got.astype(str)
    cjk = out[out.str.contains(r"[\u4e00-\u9fff]", regex=True, na=False)]
    if len(cjk):
        raise SystemExit(f"ABORT [{column}]: CJK characters survive translation: "
                         f"{sorted(set(cjk))}")
    return out


def evidence_level(cas_ok, tested):
    if cas_ok and tested:
        return "L1 (CAS + in vitro record)"
    if cas_ok and not tested:
        return "L2 (CAS only)"
    if (not cas_ok) and tested:
        return "L3 (no CAS + in vitro record)"
    return "L4 (structure record only)"


def build(df):
    df = (df.assign(_score=pd.to_numeric(df["ALT_SCORE"], errors="coerce"))
            .sort_values("_score", ascending=False, kind="mergesort")
            .drop(columns="_score").reset_index(drop=True))
    d = pd.DataFrame()
    d["No."] = np.arange(1, len(df) + 1)

    # --- Identity -------------------------------------------------------
    d["DTXSID"] = df["dtxsid"]
    d["CASRN"] = df["casrn"].fillna("")
    d["InChIKey"] = df["inchikey"]
    d["PreferredName"] = df["preferred_name"]
    d["MolecularFormula"] = df["mol_formula"]
    d["MonoisotopicMass_Da"] = pd.to_numeric(df["monoisotopic_mass"], errors="coerce").round(5)
    d["SMILES"] = df["smiles_raw"]
    d["SMILES_MS_ready"] = df["smiles_ms_ready"].fillna("")

    # --- Chemical annotation --------------------------------------------
    d["PFASClass"] = translate(df["chem_class"], CLASS_MAP, "PFASClass")
    d["FunctionalGroup"] = translate(df["fg"], FG_MAP, "FunctionalGroup")
    d["AltFamily"] = translate(df["alt_family"], FAMILY_MAP, "AltFamily")
    d["ChemicalSpace"] = translate(df["space"], SPACE_MAP, "ChemicalSpace")
    d["nCF2"] = pd.to_numeric(df["n_CF2"], errors="coerce").astype("Int64")
    d["nEtherO"] = pd.to_numeric(df["n_ether_O"], errors="coerce").astype("Int64")
    d["nRings"] = pd.to_numeric(df["n_rings"], errors="coerce").astype("Int64")

    # --- Priority scoring ------------------------------------------------
    d["ALT_SCORE"] = pd.to_numeric(df["ALT_SCORE"], errors="coerce").round(4)
    d["S_Novelty"] = pd.to_numeric(df["S_novelty"], errors="coerce").round(4)
    d["S_ReadAcross"] = pd.to_numeric(df["S_readacross"], errors="coerce").round(4)
    d["S_Concern"] = pd.to_numeric(df["S_concern"], errors="coerce").round(4)
    # 四个「持久性/迁移性相关结构标记」逐条列出（母本 Methods 2.6 声明为布尔），
    # 再给出 0–4 的命中数。列名不含 PMT：这些是结构标记，不是 PMT 评估。
    d["Marker_ShortChain"] = yn(df["Marker_短链(CF2<=5)"])
    d["Marker_Ether"] = yn(df["Marker_含醚键"])
    d["Marker_Ionisable"] = yn(df["Marker_可电离"])
    d["Marker_HighPerFrac"] = yn(df["Marker_高全氟碳占比(>=0.5)"])
    d["MarkerCount_0_4"] = pd.to_numeric(df["Marker_命中数"], errors="coerce").astype("Int64")
    d["MaxSim_KnownSubstitute"] = pd.to_numeric(df["max_sim_known_substitute"], errors="coerce").round(3)
    d["MaxSim_TrainingSet"] = pd.to_numeric(df["max_sim_train"], errors="coerce").round(3)
    d["RegrettableSubstitution"] = yn(df["regrettable_substitution"])
    d["AltTier"] = translate(df["alt_tier"], TIER_MAP, "AltTier")

    # --- Evidence --------------------------------------------------------
    cas_ok = df["casrn"].fillna("").str.match(r"^\d{2,7}-\d{2}-\d$")
    ta = pd.to_numeric(df["total_assays"], errors="coerce")
    aa = pd.to_numeric(df["active_assays"], errors="coerce")
    tested = ta.notna()
    active = (aa.fillna(0) > 0)
    d["Evidence_CAS_Registered"] = np.where(cas_ok, "Yes", "No")
    d["Evidence_InVitro_Tested"] = np.where(tested, "Yes", "No")
    d["Evidence_InVitro_Active"] = np.where(active, "Yes", "No")
    d["Evidence_Level"] = [evidence_level(c, t) for c, t in zip(cas_ok, tested)]

    # --- Model contrast ---------------------------------------------------
    d["Model_p_max"] = pd.to_numeric(df["p_max"], errors="coerce").round(4)
    d["Model_InDomain_Tanimoto"] = yn(df["in_domain_majority"])
    d["Model_InDomain_Consensus"] = yn(df["ad_consensus"])
    d["Model_Usable"] = yn(df["model_usable"])
    return d


# ---------------------------------------------------------------- README
def readme_sheet(n_all, n_s1, n_s2, overlap, stats):
    R = [
        ("SECTION 1 - OVERVIEW", ""),
        ("Title", "Priority list of emerging PFAS replacement candidates derived by a cheminformatics "
                  "read-across workflow, with identity-verification and evidence-grading columns."),
        ("Structure source",
         "US EPA CompTox Chemicals Dashboard PFASSTRUCT V6 suspect list (January 2026 release; "
         "21,028 records; 19,767 unique structures; accessed October 2026)."),
        ("Candidate substrate",
         "Chemical space B of the inventory (PFAS replacement chemical space), n = 2,468 structures; "
         "after exclusion of training-set members, known/marketed substitutes and near-duplicates, "
         "n = 2,390 candidates."),
        ("Scoring",
         "ALT_SCORE = 0.35 x novelty (1 - maximum similarity to the bioactivity panel) "
         "+ 0.20 x read-across score + 0.45 x structural concern, where structural concern = "
         "0.5 x persistence marker (perfluorinated carbon fraction) + 0.5 x mobility marker "
         "(chain length and ether-oxygen count). The structural-concern term aggregates "
         "persistence- and mobility-related structural markers and contains NO toxicity term: "
         "ALT_SCORE ranks structural concern and is not a persistence, mobility and toxicity (PMT) "
         "assessment, and cannot support one."),
        ("Important caveat",
         "PFASSTRUCT V6 is a curated SUSPECT LIST, not a census of substances demonstrated to be "
         "produced or released. A candidate's rank expresses structural concern combined with "
         "inventory membership - not demonstrated occurrence, production volume or environmental "
         "detection."),
        ("", ""),
        ("SECTION 2 - TABLE CONTENTS", ""),
        ("Table S1", f"Tier I priority candidates (ALT tier I), n = {n_s1}."),
        ("Table S2", f"Candidates carrying a potential regrettable-substitution flag, n = {n_s2}. "
                     f"The flag requires all three of: maximum similarity to a marketed substitute "
                     f"(GenX/HFPO-DA, ADONA, F-53B, 6:2 FTS, PFMOAA) >= 0.45; maximum similarity to "
                     f"any panel substance < 0.60; structural concern >= 0.60. "
                     f"Overlap with Table S1: n = {overlap}."),
        ("Table S3", f"Complete candidate list (tiers I-III), n = {n_all}."),
        ("", ""),
        ("SECTION 3 - COLUMN DEFINITIONS", ""),
        ("No.", "Rank within the table, ordered by descending ALT_SCORE."),
        ("DTXSID", "EPA DSSTox substance identifier (unique, persistent; 100% coverage)."),
        ("CASRN", "CAS Registry Number; blank where the inventory records no registered number."),
        ("InChIKey", "Structural hash - the recommended cross-database cross-reference key "
                     "(100% coverage; preferred over PubChem CID, which can change)."),
        ("PreferredName / MolecularFormula / MonoisotopicMass_Da",
         "As recorded in the inventory; monoisotopic mass is provided for exact-mass "
         "suspect-screening queries."),
        ("SMILES", "Full structure string as recorded in the inventory (includes salt/counter-ion "
                   "form where applicable)."),
        ("SMILES_MS_ready", "Counter-ion-stripped, neutralised SMILES suitable for mass-spectrometry "
                            "workflows; blank where not assignable."),
        ("PFASClass / FunctionalGroup / AltFamily",
         "Chemical class, functional group and alternative-family assignment used in the workflow."),
        ("ChemicalSpace", "Region of the chemical-space partition (A, classical long-chain PFAA; "
                          "B, replacement space; C, fluorinated aromatics; D, other; E, fallback)."),
        ("nCF2 / nEtherO / nRings", "Counts of CF2 units, ether oxygens and rings."),
        ("ALT_SCORE", "Composite priority score (see Scoring above)."),
        ("S_Novelty / S_ReadAcross / S_Concern", "Component scores entering ALT_SCORE (all 0-1)."),
        ("Marker_ShortChain / Marker_Ether / Marker_Ionisable / Marker_HighPerFrac",
         "The four persistence- and mobility-related structural markers, recorded independently of "
         "the score and reported as Yes/No: short-chain (<=5 perfluorinated carbons), "
         "ether-containing, ionisable, and high perfluorinated carbon fraction (>=0.5). They are "
         "structural markers, not a persistence, mobility and toxicity (PMT) assessment."),
        ("MarkerCount_0_4", "Number of the four structural markers above that are satisfied (0-4)."),
        ("MaxSim_KnownSubstitute", "Maximum Tanimoto similarity to a known or marketed PFAS substitute."),
        ("MaxSim_TrainingSet", "Maximum Tanimoto similarity to any training-set substance; low values "
                               "indicate that the candidate is structurally unseen by the model."),
        ("RegrettableSubstitution", "Yes where the candidate closely resembles a substitute already in "
                                    "use (structural analogue of a regrettable substitution)."),
        ("AltTier", "Priority tier: I (priority), II (watch), III (reserve)."),
        ("", ""),
        ("SECTION 4 - EVIDENCE GRADING (identity verification)", ""),
        ("Evidence_CAS_Registered", "Yes where a formal CAS Registry Number is recorded "
                                    "(fraction in this list: %s)." % stats["cas"]),
        ("Evidence_InVitro_Tested", "Yes where any ToxCast/Tox21 in vitro assay record exists "
                                    "(fraction: %s)." % stats["tested"]),
        ("Evidence_InVitro_Active", "Yes where at least one assay was reported active."),
        ("Evidence_Level", "L1 = CAS number + in vitro record;  L2 = CAS number only;  "
                           "L3 = no CAS number, in vitro record present;  L4 = structure record only."),
        ("Level distribution", "L1 = %d; L2 = %d; L3 = %d; L4 = %d (Table S3)."
         % (stats["L1"], stats["L2"], stats["L3"], stats["L4"])),
        ("Interpretation",
         "Evidence_Level grades the verifiability of a structure's identity in public registries. "
         "It does NOT indicate toxicological concern, production volume or environmental relevance."),
        ("", ""),
        ("SECTION 5 - DIMENSIONS NOT VERIFIED IN THIS DATASET", ""),
        ("Literature reports", "Not assessed. Suggested route: search each InChIKey/CASRN in Scopus, "
                               "Web of Science or PubMed; record the number of primary studies."),
        ("Patents", "Not assessed. Suggested route: search CASRN/InChIKey in PATENTSCOPE (WIPO) or "
                    "Espacenet; record patent families."),
        ("Environmental detection", "Not assessed. Suggested route: query the candidate m/z values "
                                    "(computed from MolecularFormula / MonoisotopicMass_Da) against "
                                    "published non-targeted screening datasets, e.g. PIMMS "
                                    "(Kudzin et al., Sci. Adv. 2026) and regional monitoring campaigns."),
        ("Production / use volume", "Not available from the inventory. Requires regulatory or industry "
                                    "data (e.g. OECD/ECHA dossiers, national chemical inventories)."),
        ("", ""),
        ("SECTION 6 - DATA SOURCES AND CITATION", ""),
        ("Structure inventory", "EPA CompTox PFASSTRUCT V6 (2026-01); cite Williams AJ et al. "
                                "Front Environ Sci 2022;10:850019 (list construction)."),
        ("Bioactivity panel (for the model contrast columns)",
         "Cheng & Ng, Environ Sci Technol 2019;53(23):13970-13980; archived at "
         "Zenodo DOI 10.5281/zenodo.22869943. Cite the specific PubChem BioAssay AID numbers used."),
        ("Read-across concept", "Figuiere R et al. Environ Sci Technol 2025;59(4):2031-2042 "
                                "(functional substitution and regrettable substitution)."),
        ("Code", "All tables are generated by code/tools/make_supplementary_tables.py from the "
                 "published pipeline; code and derived tables are released with the manuscript."),
    ]
    return pd.DataFrame(R, columns=["Item", "Description"])


def main():
    raw = pd.read_csv(SRC, dtype=str, low_memory=False)
    d = build(raw)

    # 用 d 自身已排序的行序判断，保证与输出表逐行对齐
    # 注意：必须精确匹配，str.startswith("I") 会同时命中 "II (watch)" 与 "III (reserve)"
    tier1_mask = d["AltTier"].eq("I (priority)").to_numpy()
    reg_mask = d["RegrettableSubstitution"].eq("Yes").to_numpy()

    S1 = d[tier1_mask].reset_index(drop=True)
    S2 = d[reg_mask].reset_index(drop=True)
    S3 = d.copy()
    for s in (S1, S2, S3):
        s["No."] = np.arange(1, len(s) + 1)

    overlap = int((tier1_mask & reg_mask).sum())

    # 证据等级统计（表 S3）
    lvl = S3["Evidence_Level"].str.slice(0, 2).value_counts()
    stats = {
        "cas": "%.1f%%" % (S3["Evidence_CAS_Registered"].eq("Yes").mean() * 100),
        "tested": "%.1f%%" % (S3["Evidence_InVitro_Tested"].eq("Yes").mean() * 100),
        "L1": int(lvl.get("L1", 0)), "L2": int(lvl.get("L2", 0)),
        "L3": int(lvl.get("L3", 0)), "L4": int(lvl.get("L4", 0)),
    }

    files = {
        "Table_S1_Tier_I_priority_candidates": S1,
        "Table_S2_Regrettable_substitutions": S2,
        "Table_S3_All_candidates": S3,
    }
    for name, s in files.items():
        s.to_csv(os.path.join(OUT, name + ".csv"), index=False, encoding="utf-8-sig")

    xlsx = os.path.join(OUT, "Supplementary_Tables_S1-S3.xlsx")
    # sheet names carry the row counts and are computed from the data, so a change
    # in the candidate count cannot leave a tab labelled with a stale number
    with pd.ExcelWriter(xlsx, engine="openpyxl") as xw:
        readme_sheet(len(S3), len(S1), len(S2), overlap, stats).to_excel(
            xw, sheet_name="README", index=False)
        S1.to_excel(xw, sheet_name=f"S1_TierI_{len(S1)}", index=False)
        S2.to_excel(xw, sheet_name=f"S2_Regrettable_{len(S2)}", index=False)
        S3.to_excel(xw, sheet_name=f"S3_All_{len(S3)}", index=False)

    # 格式化：列宽 / 冻结首行 / 自动筛选
    from openpyxl import load_workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    wb = load_workbook(xlsx)
    widths = {"DTXSID": 16, "CASRN": 15, "InChIKey": 28, "PreferredName": 46, "MolecularFormula": 15,
              "MonoisotopicMass_Da": 17, "SMILES": 46, "SMILES_MS_ready": 42, "PFASClass": 14,
              "FunctionalGroup": 18, "AltFamily": 40, "ChemicalSpace": 26, "nCF2": 7, "nEtherO": 9,
              "nRings": 7, "ALT_SCORE": 10, "S_Novelty": 10, "S_ReadAcross": 12, "S_Concern": 10,
              "Marker_ShortChain": 13, "Marker_Ether": 13, "Marker_Ionisable": 13,
              "Marker_HighPerFrac": 15, "MarkerCount_0_4": 14,
              "MaxSim_KnownSubstitute": 15, "MaxSim_TrainingSet": 15,
              "RegrettableSubstitution": 15, "AltTier": 12, "Evidence_CAS_Registered": 14,
              "Evidence_InVitro_Tested": 14, "Evidence_InVitro_Active": 14, "Evidence_Level": 24,
              "Model_p_max": 11, "Model_InDomain_Tanimoto": 14, "Model_InDomain_Consensus": 14,
              "Model_Usable": 11, "No.": 6}
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        if ws.title != "README":
            ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True)
        hdr = [c.value for c in ws[1]]
        for i, h in enumerate(hdr, start=1):
            w = widths.get(h, 34 if ws.title == "README" else 12)
            if ws.title == "README":
                w = 30 if i == 1 else 110
            ws.column_dimensions[get_column_letter(i)].width = w
    wb.save(xlsx)

    print("输出目录:", OUT)
    for f in sorted(os.listdir(OUT)):
        print("  ", f, os.path.getsize(os.path.join(OUT, f)), "bytes")
    print()
    print(f"S1 Tier I          : {len(S1)}")
    print(f"S2 Regrettable     : {len(S2)}  (与 S1 重叠 {overlap})")
    print(f"S3 All             : {len(S3)}")
    print("Evidence_Level 分布 L1/L2/L3/L4 : %d / %d / %d / %d" % (stats["L1"], stats["L2"], stats["L3"], stats["L4"]))
    print("CAS 登记比例: %s   有体外记录比例: %s" % (stats["cas"], stats["tested"]))


if __name__ == "__main__":
    main()
